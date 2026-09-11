"""Test fixtures.

The pitfall this works around: `database/session.py` builds its engine at
module import time from `settings`, and `config.get_settings` is lru_cached.
So the environment has to be set and the modules dropped from `sys.modules`
*before* anything imports them - which is why the imports below happen inside
the fixture rather than at the top of the file.
"""

from __future__ import annotations

import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

_APP_PACKAGES = {
    "config",
    "controller",
    "database",
    "helper",
    "llm",
    "main",
    "repository",
    "schema",
    "service",
}


def _reset_app_modules() -> None:
    for name in [m for m in sys.modules if m.split(".")[0] in _APP_PACKAGES]:
        del sys.modules[name]


@pytest.fixture()
def app_modules(tmp_path, monkeypatch):
    """A fresh app wired to a throwaway database, fully offline.

    Both network layers are disabled so the retrieval assertions are
    deterministic: no Groq key means `ask` still returns its sources with
    `llm_configured: false`, and no HF token means the dense channel is empty.
    Retrieval must work anyway - that is the point of the regression test.
    """
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'test.db'}")
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("GROQ_API_KEY", "")
    monkeypatch.setenv("HF_TOKEN", "")

    _reset_app_modules()

    import config

    config.get_settings.cache_clear()

    from database.session import SessionLocal, init_db
    from repository import document_repository
    from service import document_service, qa_service

    init_db()

    yield SimpleNamespace(
        settings=config.settings,
        SessionLocal=SessionLocal,
        repository=document_repository,
        documents=document_service,
        qa=qa_service,
    )

    _reset_app_modules()


@pytest.fixture()
def db(app_modules):
    session = app_modules.SessionLocal()
    try:
        yield session
    finally:
        session.close()


def build_docx(path: Path) -> bytes:
    """A small, realistic .docx: a title, headings and prose."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("System Design Roadmap", level=0)
    doc.add_paragraph(
        "A forty day study plan that moves from the basics of distributed "
        "systems up to complete architectures. One topic per day."
    )
    doc.add_heading("Phase 1 - Foundations", level=1)
    doc.add_paragraph(
        "Latency and throughput, the client server model, and how DNS resolves "
        "a hostname before any request is sent."
    )
    doc.add_heading("Phase 2 - Data at Scale", level=1)
    doc.add_paragraph(
        "Sharding splits a table across machines by a partition key. "
        "Replication keeps redundant copies so a single node loss is survivable."
    )
    doc.save(path)
    return path.read_bytes()


def build_pdf(path: Path) -> bytes:
    """A two page PDF: page 1 is prose with no drawings, page 2 a ruled table.

    The split is the point. Table detection only visits pages that carry
    ruling lines, so a fixture needs one page of each to show both that the
    prose page is skipped and that the table page is still read.
    """
    import pymupdf

    doc = pymupdf.open()

    page = doc.new_page()
    y = 90
    for line in (
        "Phase 1 - Foundations",
        "Latency and throughput, the client server model, and how DNS",
        "resolves a hostname before any request is sent to the origin.",
        "Sharding splits a table across machines by a partition key.",
    ):
        page.insert_text((72, y), line, fontsize=11)
        y += 18

    page = doc.new_page()
    left, top, width, height, rows, cols = 72, 90, 360, 60, 3, 3
    for row in range(rows + 1):
        page.draw_line(
            pymupdf.Point(left, top + row * height),
            pymupdf.Point(left + width, top + row * height),
        )
    for col in range(cols + 1):
        page.draw_line(
            pymupdf.Point(left + col * width / cols, top),
            pymupdf.Point(left + col * width / cols, top + rows * height),
        )
    cells = [
        ["Region", "Replicas", "Latency"],
        ["us-east", "3", "12ms"],
        ["eu-west", "2", "28ms"],
    ]
    for row, values in enumerate(cells):
        for col, value in enumerate(values):
            page.insert_text(
                (left + col * width / cols + 8, top + row * height + 22),
                value,
                fontsize=10,
            )

    doc.save(path)
    doc.close()
    return path.read_bytes()
