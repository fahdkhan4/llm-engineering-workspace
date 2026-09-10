"""Word (.docx) extraction using python-docx.

Paragraph text is read straight from the underlying XML rather than through
`Paragraph.text`, because the latter silently drops text that lives inside
hyperlinks, text boxes and shapes.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile

from docx import Document as open_docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from database.models import BlockType
from helper.extraction_models import (
    ExtractedBlock,
    ExtractedDocument,
    ExtractedTable,
    rows_to_markdown,
)

logger = logging.getLogger(__name__)

_WHITESPACE_RE = re.compile(r"\s+")
_HEADING_RE = re.compile(r"^heading\s*(\d+)$", re.IGNORECASE)
_APP_PROPS = "docProps/app.xml"


def extract_docx(data: bytes) -> ExtractedDocument:
    """Extract text, structure, tables and metadata from .docx bytes."""
    document = open_docx(io.BytesIO(data))
    result = ExtractedDocument(metadata=_metadata(document, data))

    header_texts, footer_texts = _headers_and_footers(document)
    for text in header_texts:
        result.blocks.append(
            ExtractedBlock(page_number=1, block_type=BlockType.HEADER, content=text)
        )

    _walk(document, result)

    for text in footer_texts:
        result.blocks.append(
            ExtractedBlock(page_number=1, block_type=BlockType.FOOTER, content=text)
        )

    result.page_count = int(result.metadata.get("pages") or 1)
    if not result.blocks:
        result.ocr_required_pages = [1]
    return result


def _walk(container, result: ExtractedDocument) -> None:
    """Emit blocks for one container (the body, or a table cell) in order."""
    for item in _iter_block_items(container):
        if isinstance(item, Table):
            _add_table(item, result)
            continue

        text = _paragraph_text(item)
        if not text:
            continue

        level = _heading_level(item)
        if level is not None:
            result.blocks.append(
                ExtractedBlock(
                    page_number=1,
                    block_type=BlockType.HEADING,
                    content=text,
                    heading_level=level,
                    meta={"style": item.style.name if item.style else None},
                )
            )
            continue

        block_type = BlockType.LIST_ITEM if _is_list_item(item) else BlockType.PARAGRAPH
        result.blocks.append(
            ExtractedBlock(
                page_number=1,
                block_type=block_type,
                content=text,
                meta={"style": item.style.name if item.style else None},
            )
        )


def _add_table(table: Table, result: ExtractedDocument) -> None:
    rows = [
        [_WHITESPACE_RE.sub(" ", cell.text).strip() for cell in row.cells]
        for row in table.rows
    ]
    rows = [row for row in rows if any(row)]
    if not rows:
        return

    index = len(result.tables)
    extracted = ExtractedTable(
        page_number=1, table_index=index, rows=rows, markdown=rows_to_markdown(rows)
    )
    result.tables.append(extracted)
    result.blocks.append(
        ExtractedBlock(
            page_number=1,
            block_type=BlockType.TABLE,
            content=extracted.markdown,
            table_index=index,
            meta={"rows": extracted.n_rows, "columns": extracted.n_cols},
        )
    )


def _iter_block_items(parent):
    """Yield paragraphs and tables in document order (python-docx keeps them apart)."""
    if isinstance(parent, DocxDocument):
        parent_element = parent.element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    else:
        raise TypeError(f"Cannot iterate over {type(parent)!r}")

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _paragraph_text(paragraph: Paragraph) -> str:
    """All text under the paragraph, including hyperlinks and text boxes."""
    parts = [node.text or "" for node in paragraph._p.xpath(".//w:t")]
    return _WHITESPACE_RE.sub(" ", "".join(parts)).strip()


def _heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style.name if paragraph.style else ""
    match = _HEADING_RE.match(style or "")
    if match:
        return min(int(match.group(1)), 9)
    if style in ("Title",):
        return 1
    if style in ("Subtitle",):
        return 2

    # Styles can be renamed; the outline level is the reliable fallback.
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None:
            value = outline.get(qn("w:val"))
            if value is not None and value.isdigit() and int(value) < 9:
                return int(value) + 1
    return None


def _is_list_item(paragraph: Paragraph) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.numPr is not None:
        return True
    style = (paragraph.style.name if paragraph.style else "") or ""
    return "list" in style.lower()


def _headers_and_footers(document: DocxDocument) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    footers: list[str] = []
    for section in document.sections:
        for part, bucket in (
            (section.header, headers),
            (section.first_page_header, headers),
            (section.footer, footers),
            (section.first_page_footer, footers),
        ):
            try:
                for paragraph in part.paragraphs:
                    text = _paragraph_text(paragraph)
                    if text and text not in bucket:
                        bucket.append(text)
            except Exception:
                continue
    return headers, footers


def _metadata(document: DocxDocument, data: bytes) -> dict:
    props = document.core_properties
    meta = {
        "title": props.title,
        "author": props.author,
        "subject": props.subject,
        "keywords": props.keywords,
        "category": props.category,
        "comments": props.comments,
        "last_modified_by": props.last_modified_by,
        "revision": props.revision,
        "creationDate": props.created.isoformat() if props.created else None,
        "modDate": props.modified.isoformat() if props.modified else None,
    }
    meta.update(_app_properties(data))
    return {key: value for key, value in meta.items() if value}


def _app_properties(data: bytes) -> dict:
    """Page/word counts Word records in docProps/app.xml (not exposed by python-docx)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read(_APP_PROPS).decode("utf-8", errors="ignore")
    except (KeyError, zipfile.BadZipFile, OSError):
        return {}

    found = {}
    for tag in ("Pages", "Words", "Characters", "Company", "Application"):
        match = re.search(rf"<{tag}>([^<]*)</{tag}>", xml)
        if match and match.group(1).strip():
            found[tag.lower()] = match.group(1).strip()
    return found
